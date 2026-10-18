import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top=True, bottom=True, left=False, right=False):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for b_name, b_val in borders.items():
        node = OxmlElement(f'w:{b_name}')
        if b_val:
            node.set(qn('w:val'), 'single')
            node.set(qn('w:sz'), '4')
            node.set(qn('w:space'), '0')
            node.set(qn('w:color'), '000000')
        else:
            node.set(qn('w:val'), 'none')
        tcBorders.append(node)
    tcPr.append(tcBorders)

def add_code_block(doc, code_str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = Pt(14)
    run = p.add_run(code_str)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 30, 30)

def generate_theory_docx(output_path):
    doc = docx.Document()
    
    # Page margins
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)
        
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(12)
    title_run = title_p.add_run('2.1.1. ÔN TẬP LÝ THUYẾT: CÂY QUYẾT ĐỊNH VÀ RỪNG NGẪU NHIÊN')
    title_run.font.name = 'Times New Roman'
    title_run.font.size = Pt(14)
    title_run.bold = True
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_q_header(num):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(f'Câu {num}.')
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.bold = True
        return p

    def add_body(text, bold_prefix='', space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(11)
            r_pre.bold = True
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        return p

    # Câu 1
    add_q_header(1)
    add_body('CRISP-DM là quy trình tiêu chuẩn liên ngành phục vụ khai phá dữ liệu gồm 6 giai đoạn tuần tự và có tính chu kỳ:')
    add_body('1. Thấu hiểu nghiệp vụ: Xác định mục tiêu bài toán thực tế, yêu cầu nghiệp vụ và tiêu chuẩn đánh giá thành công.')
    add_body('2. Thấu hiểu dữ liệu: Thu thập dữ liệu ban đầu, khảo sát cấu trúc, kiểm tra tính toàn vẹn và phát hiện các vấn đề chất lượng.')
    add_body('3. Chuẩn bị dữ liệu: Làm sạch dữ liệu, xử lý giá trị khuyết thiếu, mã hóa biến phân loại và trích xuất các đặc trưng phù hợp.')
    add_body('4. Xây dựng mô hình: Lựa chọn kỹ thuật học máy thích hợp, chia tập huấn luyện và kiểm tra, thiết lập tham số và tiến hành huấn luyện.')
    add_body('5. Đánh giá: So sánh hiệu năng các mô hình dựa trên thước đo kỹ thuật và kiểm tra mức độ đáp ứng mục tiêu nghiệp vụ ban đầu.')
    add_body('6. Triển khai: Đóng gói mô hình thành dịch vụ phần mềm hoặc báo cáo kết quả để phục vụ người dùng cuối.')
    add_body('SEMMA do viện SAS phát triển, tập trung vào các bước kỹ thuật trong việc thao tác và khai phá dữ liệu gồm 5 bước:')
    add_body('1. Sample: Trích chọn mẫu dữ liệu đại diện từ kho dữ liệu lớn để đảm bảo tốc độ xử lý.')
    add_body('2. Explore: Khám phá trực quan và thống kê để nhận diện xu hướng, phân phối và điểm dị biệt.')
    add_body('3. Modify: Biến đổi dữ liệu, xử lý ngoại lai, tạo biến mới và chuẩn hóa thang đo.')
    add_body('4. Model: Áp dụng các thuật toán học máy để khớp dữ liệu.')
    add_body('5. Assess: Đánh giá độ chính xác và tính hữu dụng của các mô hình đã xây dựng.')
    add_body('So sánh sự khác biệt giữa hai quy trình:')
    
    # Table 1
    table1 = doc.add_table(rows=4, cols=3)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ['Tiêu chí so sánh', 'Quy trình CRISP-DM', 'Quy trình SEMMA']
    data1 = [
        ['Phạm vi bao phủ', 'Bao quát toàn bộ vòng đời dự án từ bài toán thực tế đến vận hành hệ thống', 'Tập trung chủ yếu vào các bước kỹ thuật phân tích và mô hình hóa dữ liệu'],
        ['Giai đoạn định hướng', 'Có giai đoạn thấu hiểu nghiệp vụ độc lập ở đầu quy trình', 'Không có giai đoạn nghiệp vụ riêng biệt, bắt đầu ngay từ việc chọn mẫu dữ liệu'],
        ['Giai đoạn kết thúc', 'Có bước đánh giá nghiệp vụ và triển khai sản phẩm thực tế', 'Kết thúc ở bước đánh giá độ chính xác của mô hình']
    ]
    for j, h in enumerate(headers):
        cell = table1.cell(0, j)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.bold = True
        set_cell_margins(cell, 80, 80, 100, 100)
        set_cell_borders(cell, top=True, bottom=True, left=False, right=False)
    for i, row in enumerate(data1):
        for j, val in enumerate(row):
            cell = table1.cell(i + 1, j)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            set_cell_margins(cell, 80, 80, 100, 100)
            set_cell_borders(cell, top=(i == 0), bottom=(i == len(data1) - 1), left=False, right=False)

    # Câu 2
    add_q_header(2)
    add_body('Cây quyết định là thuật toán học có giám sát hoạt động theo nguyên lý chia để trị. Mô hình liên tục phân chia tập dữ liệu thành các tập con thuần khiết hơn thông qua các câu hỏi điều kiện trên từng thuộc tính.')
    add_body('Cấu trúc cây quyết định gồm ba thành phần cơ bản:')
    add_body('- Nút gốc: Nút nằm ở vị trí cao nhất trên cây, chứa toàn bộ dữ liệu huấn luyện ban đầu. Tại đây, thuật toán tính toán độ thuần khiết của tất cả thuộc tính và chọn thuộc tính tốt nhất để thực hiện phép chia đầu tiên.')
    add_body('- Nút nhánh: Các nút trung gian nhận dữ liệu từ nút cha và tiếp tục áp dụng một phép kiểm tra điều kiện trên một thuộc tính cụ thể để chia dữ liệu thành hai hoặc nhiều nhánh con.')
    add_body('- Nút lá: Các nút kết thúc nằm ở cuối mỗi nhánh, không tiếp tục phân chia. Mỗi nút lá chứa nhãn dự đoán cho bài toán phân loại hoặc một giá trị số trung bình cho bài toán hồi quy.')
    add_body('Cơ chế dự đoán:')
    add_body('Khi có một mẫu dữ liệu mới cần dự đoán, mẫu sẽ xuất phát từ nút gốc. Giá trị của thuộc tính tương ứng trên mẫu sẽ được so sánh với ngưỡng điều kiện tại nút. Nếu thỏa mãn điều kiện, mẫu đi theo nhánh đúng; nếu không, mẫu đi theo nhánh sai. Quá trình kiểm tra lặp lại tại các nút nhánh tiếp theo cho đến khi mẫu chạm tới nút lá. Nhãn của nút lá đó chính là kết quả dự đoán của mô hình.')

    # Câu 3
    add_q_header(3)
    add_body('Các tiêu chí phân tách được dùng để đo lường mức độ tinh khiết của một nút và lựa chọn điểm cắt tối ưu nhất.')
    add_body('Chỉ số Gini:')
    add_body('Công thức tính: Gini = 1 - tổng bình phương xác suất các lớp i từ 1 đến C.')
    add_body('Ý nghĩa: Gini đo xác suất một phần tử bất kỳ bị gán sai nhãn nếu được gán nhãn ngẫu nhiên theo phân bố xác suất trong nút. Giá trị Gini nằm trong khoảng từ 0 đến 1 - 1/C. Nút thuần khiết chỉ chứa mẫu của một lớp duy nhất sẽ có Gini bằng 0. Nút càng hỗn tạp thì Gini càng cao.')
    add_body('Entropy:')
    add_body('Công thức tính: Entropy = - tổng xác suất lớp i nhân logarit cơ số 2 của xác suất lớp i từ 1 đến C.')
    add_body('Ý nghĩa: Entropy xuất phát từ lý thuyết thông tin, đo lường mức độ hỗn loạn của dữ liệu. Nếu nút chỉ chứa một lớp, Entropy bằng 0. Nếu các lớp phân bố đều nhau, Entropy đạt giá trị cực đại bằng logarit cơ số 2 của C.')
    add_body('Information Gain:')
    add_body('Công thức tính: Information Gain = Entropy của tập dữ liệu trước khi chia - tổng tỷ trọng kích thước nút con nhân Entropy nút con theo từng giá trị của thuộc tính A.')
    add_body('Ý nghĩa: Đo lường lượng thông tin thu được, hay mức độ giảm độ hỗn loạn của dữ liệu sau khi chia tập dữ liệu theo thuộc tính A. Thuộc tính nào mang lại Information Gain cao nhất sẽ được chọn làm điều kiện phân tách tại nút đó.')
    add_body('Sự khác nhau cốt lõi:')
    add_body('Gini sử dụng phép tính bình phương đơn giản nên tốc độ tính toán nhanh hơn nhiều so với Entropy vốn yêu cầu phép tính logarit. Về mặt phân tách, Gini và Entropy cho kết quả tương đương nhau trong đại đa số trường hợp thực tế.')

    # Câu 4
    add_q_header(4)
    add_body('Random Forest là mô hình học kết hợp xây dựng từ một tập hợp gồm nhiều cây quyết định độc lập.')
    add_body('Khác biệt giữa Random Forest và cây quyết định đơn lẻ:')
    add_body('- Cây quyết định đơn lẻ sử dụng toàn bộ tập dữ liệu huấn luyện và xét toàn bộ các đặc trưng tại mỗi điểm phân tách. Cấu trúc này làm cây rất nhạy cảm với dữ liệu huấn luyện, dễ ghi nhớ cả nhiễu và dẫn đến hiện tượng quá khớp.')
    add_body('- Random Forest tạo ra nhiều cây con độc lập bằng kỹ thuật lấy mẫu ngẫu nhiên có hoàn lại từ tập huấn luyện ban đầu. Đồng thời, tại mỗi nút phân tách của từng cây, thuật toán chỉ chọn ngẫu nhiên một tập con các đặc trưng thay vì xét toàn bộ đặc trưng.')
    add_body('Lý do Random Forest có hiệu suất cao hơn và giảm quá khớp:')
    add_body('Mỗi cây con trong rừng có độ chệch thấp nhưng phương sai cao. Do dữ liệu huấn luyện và tập đặc trưng được lấy ngẫu nhiên, các cây con này ít tương quan với nhau. Khi kết hợp kết quả bằng cơ chế bỏ phiếu đa số, phương sai tổng thể của mô hình giảm mạnh theo quy luật số lớn trong khi độ chệch vẫn giữ ở mức thấp. Nhờ vậy, sai số dự đoán trên dữ liệu mới giảm đáng kể và hiện tượng quá khớp được kiểm soát hiệu quả.')

    # Câu 5
    add_q_header(5)
    add_body('Ưu điểm và hạn chế của Decision Tree:')
    add_body('- Ưu điểm: Cấu trúc trực quan, dễ giải thích quy tắc rẽ nhánh, không cần chuẩn hóa độ lớn của các biến số, tự động xử lý được cả dữ liệu dạng số và dạng phân loại.')
    add_body('- Hạn chế: Rất dễ bị quá khớp nếu không giới hạn độ sâu; không ổn định vì một thay đổi nhỏ trong dữ liệu huấn luyện có thể làm thay đổi hoàn toàn cấu trúc cây; ranh giới phân tách chỉ là các đường song song với trục tọa độ.')
    add_body('- Hoạt động kém hiệu quả khi: Dữ liệu có nhiều biến liên tục có tương quan phức tạp không song song với trục tọa độ, dữ liệu mất cân bằng nghiêm trọng hoặc khi tập dữ liệu chứa nhiều mẫu nhiễu.')
    add_body('Ưu điểm và hạn chế của Random Forest:')
    add_body('- Ưu điểm: Hiệu năng phân loại vượt trội so với cây đơn lẻ, kháng hiện tượng quá khớp rất tốt, hoạt động ổn định trên dữ liệu nhiều chiều và cho phép trích xuất độ quan trọng của từng đặc trưng.')
    add_body('- Hạn chế: Thời gian huấn luyện và dung lượng bộ nhớ lớn; khó diễn giải trực tiếp quy tắc ra quyết định vì mô hình là sự kết hợp của hàng trăm cây khác nhau; tốc độ dự đoán thời gian thực chậm hơn cây đơn lẻ.')

    # Câu 6
    add_q_header(6)
    add_body('Đoạn mã Python sử dụng Scikit-learn xây dựng cây quyết định gồm 5 bước cụ thể:')
    code_dt = (
        "from sklearn.datasets import load_iris\n"
        "from sklearn.model_selection import train_test_split\n"
        "from sklearn.tree import DecisionTreeClassifier\n"
        "from sklearn.metrics import accuracy_score, classification_report\n\n"
        "# Buoc 1: Nap du lieu\n"
        "iris = load_iris()\n"
        "X, y = iris.data, iris.target\n\n"
        "# Buoc 2: Chia tap du lieu 80% huan luyen va 20% kiem tra\n"
        "X_train, X_test, y_train, y_test = train_test_split(\n"
        "    X, y, test_size=0.2, random_state=42, stratify=y\n"
        ")\n\n"
        "# Buoc 3: Khoi tao mo hinh cay quyet dinh voi do sau toi da bang 3\n"
        "model = DecisionTreeClassifier(criterion='gini', max_depth=3, random_state=42)\n\n"
        "# Buoc 4: Huan luyen mo hinh tren tap huan luyen\n"
        "model.fit(X_train, y_train)\n\n"
        "# Buoc 5: Du doan tren tap kiem tra va danh gia ket qua\n"
        "y_pred = model.predict(X_test)\n"
        "print('Accuracy:', accuracy_score(y_test, y_pred))\n"
        "print(classification_report(y_test, y_pred, target_names=iris.target_names))\n"
    )
    add_code_block(doc, code_dt)
    add_body('Mô tả quy trình:')
    add_body('Bước 1 tải tập dữ liệu 150 mẫu gồm 4 đặc trưng và 3 loài hoa. Bước 2 phân chia ngẫu nhiên giữ nguyên tỷ lệ nhãn lớp giữa hai tập. Bước 3 tạo đối tượng cây với tiêu chí Gini và khống chế độ sâu bằng 3 để tránh quá khớp. Bước 4 khớp các điểm cắt điều kiện trên tập huấn luyện. Bước 5 đánh giá độ chính xác và xuất báo cáo phân loại chi tiết.')

    # Câu 7
    add_q_header(7)
    add_body('Cách triển khai Random Forest trong Python:')
    code_rf = (
        "from sklearn.ensemble import RandomForestClassifier\n\n"
        "rf_model = RandomForestClassifier(\n"
        "    n_estimators=100,\n"
        "    max_depth=5,\n"
        "    max_features='sqrt',\n"
        "    min_samples_split=5,\n"
        "    min_samples_leaf=2,\n"
        "    random_state=42\n"
        ")\n"
        "rf_model.fit(X_train, y_train)\n"
    )
    add_code_block(doc, code_rf)
    add_body('Ý nghĩa của 5 siêu tham số chính:')
    add_body('- n_estimators: Tổng số cây quyết định được tạo trong rừng ngẫu nhiên. Số lượng cây càng nhiều thì kết quả phân loại càng ổn định và giảm phương sai, tuy nhiên thời gian tính toán sẽ tăng tỉ lệ thuận.')
    add_body('- max_depth: Chiều sâu tối đa cho phép của mỗi cây con. Giới hạn độ sâu giúp ngăn từng cây con học quá chi tiết vào dữ liệu huấn luyện, từ đó kiểm soát quá khớp.')
    add_body('- max_features: Số lượng đặc trưng ngẫu nhiên tối đa được chọn để tìm điểm phân tách tại mỗi nút. Thiết lập căn bậc hai của tổng số đặc trưng giúp các cây con ít tương quan với nhau hơn.')
    add_body('- min_samples_split: Số lượng mẫu dữ liệu tối thiểu cần thiết tại một nút để nút đó được phép tiếp tục phân chia. Giá trị lớn giúp ngăn việc tạo ra các nhánh chỉ chứa quá ít mẫu.')
    add_body('- min_samples_leaf: Số lượng mẫu tối thiểu bắt buộc phải có tại mỗi nút lá. Tham số này bảo đảm mỗi quyết định phân loại ở nút kết thúc đều dựa trên một số lượng mẫu đủ tin cậy.')

    # Câu 8
    add_q_header(8)
    add_body('Trong Random Forest, độ quan trọng của đặc trưng được tính dựa trên mức giảm độ thuần khiết trung bình qua tất cả các cây trong rừng khi một đặc trưng được chọn để chia nút. Scikit-learn lưu trữ giá trị này trong thuộc tính feature_importances_.')
    add_body('Đoạn mã trích xuất và vẽ biểu đồ thanh ngang:')
    code_fi = (
        "import pandas as pd\n"
        "import matplotlib.pyplot as plt\n\n"
        "# Trich xuat do quan trong cua tung dac trung\n"
        "importances = pd.Series(rf_model.feature_importances_, index=iris.feature_names)\n"
        "importances_sorted = importances.sort_values(ascending=True)\n\n"
        "# Ve bieu do thanh ngang\n"
        "plt.figure(figsize=(8, 4))\n"
        "importances_sorted.plot(kind='barh', color='steelblue')\n"
        "plt.xlabel('Do quan trong')\n"
        "plt.ylabel('Dac trung')\n"
        "plt.title('Muc do quan trong cua dac trung trong Random Forest')\n"
        "plt.tight_layout()\n"
        "plt.show()\n"
    )
    add_code_block(doc, code_fi)
    add_body('Biểu đồ sắp xếp tăng dần giúp người phân tích nhận biết ngay đặc trưng nào đóng góp lớn nhất vào khả năng phân loại của mô hình.')

    # Câu 9
    add_q_header(9)
    add_body('So sánh hai phương pháp điều chỉnh siêu tham số:')
    add_body('- GridSearchCV duyệt vét cạn qua tất cả các tổ hợp có thể có trong lưới tham số định sẵn. Với mỗi tổ hợp, phương pháp thực hiện kiểm định chéo K-fold để tính điểm đánh giá trung bình. Phương pháp này bảo đảm không bỏ sót cấu hình tối ưu trong lưới nhưng tốn nhiều thời gian khi số tham số lớn.')
    add_body('- RandomizedSearchCV chọn ngẫu nhiên một số lượng cấu hình cố định từ không gian tham số thay vì duyệt toàn bộ. Cách tiếp cận này tiết kiệm đáng kể thời gian tính toán và phù hợp khi cần khảo sát không gian tham số rộng.')
    add_body('Đoạn mã Python triển khai cả hai phương pháp:')
    code_tune = (
        "from sklearn.model_selection import GridSearchCV, RandomizedSearchCV\n\n"
        "param_grid = {\n"
        "    'n_estimators': [50, 100, 150],\n"
        "    'max_depth': [3, 5, 7],\n"
        "    'min_samples_split': [2, 5]\n"
        "}\n\n"
        "# Dieu chinh vet can bang GridSearchCV voi 5-fold cross validation\n"
        "grid_search = GridSearchCV(\n"
        "    estimator=RandomForestClassifier(random_state=42),\n"
        "    param_grid=param_grid,\n"
        "    cv=5,\n"
        "    scoring='accuracy'\n"
        ")\n"
        "grid_search.fit(X_train, y_train)\n"
        "print('Best params GridSearchCV:', grid_search.best_params_)\n\n"
        "# Dieu chinh ngau nhien bang RandomizedSearchCV thu 10 to hop\n"
        "random_search = RandomizedSearchCV(\n"
        "    estimator=RandomForestClassifier(random_state=42),\n"
        "    param_distributions=param_grid,\n"
        "    n_iter=10,\n"
        "    cv=5,\n"
        "    scoring='accuracy',\n"
        "    random_state=42\n"
        ")\n"
        "random_search.fit(X_train, y_train)\n"
        "print('Best params RandomizedSearchCV:', random_search.best_params_)\n"
    )
    add_code_block(doc, code_tune)

    doc.save(output_path)
    print('Saved theory document to:', output_path)

if __name__ == '__main__':
    generate_theory_docx('Lab05_ThucHanhBaiToanPhanLop/Lab_2_1_1_On_tap_ly_thuyet_Decision_Tree_Random_Forest.docx')
    generate_theory_docx('Lab05_ThucHanhBaiToanPhanLop/code/Reports/Lythuyet_2_1_DecisionTree_RandomForest.docx')
