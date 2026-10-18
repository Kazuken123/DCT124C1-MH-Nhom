import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top=True, bottom=True, left=False, right=False):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for b_name, b_val in borders.items():
        node = OxmlElement(f'w:{b_name}')
        if b_val:
            node.set(qn('w:val'), 'single')
            node.set(qn('w:sz'), '4')
            node.set(qn('w:space'), '0')
            node.set(qn('w:color'), '000000')
        else:
            node.set(qn('w:val'), 'none')
        tcBorders.append(node)
    tcPr.append(tcBorders)

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    r.italic = True
    return p

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12.5)
    r.bold = True
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)
    r.bold = True
    return p

def add_body(doc, text, bold_prefix='', space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(11)
        r_pre.bold = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    return p

def add_image_safe(doc, img_path, width_in=5.8, caption=''):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(img_path, width=Inches(width_in))
        if caption:
            add_caption(doc, caption)
    else:
        print(f"Warning: Image not found {img_path}")

def format_table_bw(table, headers, data, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.bold = True
        set_cell_margins(cell, 80, 80, 100, 100)
        set_cell_borders(cell, top=True, bottom=True, left=False, right=False)
        if col_widths and j < len(col_widths):
            cell.width = Inches(col_widths[j])

    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            if j == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(val))
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            set_cell_margins(cell, 80, 80, 100, 100)
            is_last = (i == len(data) - 1)
            set_cell_borders(cell, top=(i == 0), bottom=is_last, left=False, right=False)
            if col_widths and j < len(col_widths):
                cell.width = Inches(col_widths[j])

def generate_report():
    doc = docx.Document()
    
    # Page setup
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)
        
    # COVER PAGE
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_uni.paragraph_format.space_before = Pt(0)
    p_uni.paragraph_format.space_after = Pt(2)
    r = p_uni.add_run("TRƯỜNG ĐẠI HỌC SÀI GÒN\nKHOA CÔNG NGHỆ THÔNG TIN")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.bold = True

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(36)
    p_space.paragraph_format.space_after = Pt(12)
    p_space.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_space.add_run("BÁO CÁO THỰC HÀNH MÔN HỌC MÁY HỌC\nLỚP: DCT124C1")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(14)
    r_sub.bold = True

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(18)
    p_title.paragraph_format.space_after = Pt(36)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("BÀI TOÁN PHÂN LỚP DỮ LIỆU\nCÂY QUYẾT ĐỊNH, RỪNG NGẪU NHIÊN, SUPPORT VECTOR MACHINE VÀ NAIVE BAYES")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(16)
    r_t.bold = True

    # Info table
    info_table = doc.add_table(rows=6, cols=3)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_headers = ["Họ và tên", "MSSV", "Vai trò và Nhiệm vụ phụ trách"]
    info_data = [
        ["Lê Văn Hiếu", "3124411091", "Nhóm trưởng - Phụ trách Mục 2.1 Cây quyết định và Rừng ngẫu nhiên, tổng hợp báo cáo"],
        ["Văn Nguyễn Thành Đạt", "3124411067", "Thành viên - Phụ trách Mục 2.2 Support Vector Machine"],
        ["Dương Gia Phát", "3124411212", "Thành viên - Phụ trách Mục 2.3 Naive Bayes"],
        ["Cái Trần Minh Tiến", "3124411309", "Thành viên - Phụ trách Phần 1 Khám phá dữ liệu EDA"],
        ["Giảng viên hướng dẫn:", "TS. Đỗ Như Tài", "Bộ môn Khoa học Máy tính - Khoa Công nghệ Thông tin"]
    ]
    format_table_bw(info_table, info_headers, info_data, [1.6, 1.2, 3.8])

    p_date = doc.add_paragraph()
    p_date.paragraph_format.space_before = Pt(48)
    p_date.paragraph_format.space_after = Pt(0)
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_d = p_date.add_run("Thành phố Hồ Chí Minh, Năm 2026")
    r_d.font.name = "Times New Roman"
    r_d.font.size = Pt(11)
    
    doc.add_page_break()

    # MỤC LỤC
    add_heading_1(doc, "MỤC LỤC TỔNG QUAN")
    toc_items = [
        ("PHẦN 1: KHÁM PHÁ VÀ TRỰC QUAN HÓA DỮ LIỆU", "1"),
        ("  1.1 Khảo sát tập dữ liệu chim cánh cụt Penguins", "1"),
        ("  1.2 Khảo sát tập dữ liệu giá nhà ở Housing Amsterdam", "2"),
        ("  1.3 Khảo sát tập dữ liệu chiến dịch tiếp thị Marketing Campaign", "2"),
        ("PHẦN 2: THỰC HÀNH CÁC GIẢI THUẬT PHÂN LỚP", "3"),
        ("  MỤC 2.1: CÂY QUYẾT ĐỊNH VÀ RỪNG NGẪU NHIÊN", "3"),
        ("    2.1.1 Ôn tập lý thuyết", "3"),
        ("    2.1.2 Bài làm mẫu: Dự đoán khả năng vỡ nợ thẻ tín dụng", "6"),
        ("    2.1.3 Bài tập thực hành 1: Dự đoán khả năng sống sót tàu Titanic", "8"),
        ("    2.1.4 Bài tập thực hành 2: Dự đoán bệnh tiểu đường Pima Indians", "10"),
        ("    2.1.5 Bảng so sánh trực diện và nhận xét hiện tượng quá khớp", "12"),
        ("  MỤC 2.2: BÀI TOÁN PHÂN LỚP VỚI SUPPORT VECTOR MACHINE", "13"),
        ("    2.2.1 Ôn tập lý thuyết SVM", "13"),
        ("    2.2.2 Thực nghiệm phân loại hoa Iris", "15"),
        ("    2.2.3 Thực nghiệm nhận dạng chữ số viết tay Digits", "16"),
        ("    2.2.4 Thực nghiệm dự đoán bệnh tiểu đường Diabetes", "17"),
        ("    2.2.5 Thực nghiệm phân loại tình trạng động vật Animal Condition", "18"),
        ("  MỤC 2.3: BÀI TOÁN PHÂN LỚP VỚI BAYES NGÂY THƠ", "19"),
        ("    2.3.1 Ôn tập lý thuyết Naive Bayes", "19"),
        ("    2.3.2 Bài làm mẫu: Phân loại Email Spam", "21"),
        ("    2.3.3 Bài tập thực hành 1: Dự đoán hành vi mua hàng Customer Behaviour", "22"),
        ("    2.3.4 Bài tập thực hành 2: Phân loại nấm Mushroom", "23"),
        ("PHẦN 3: TỔNG KẾT VÀ SO SÁNH CÁC NHÓM MÔ HÌNH", "24")
    ]
    for item, page in toc_items:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(1)
        p_t.paragraph_format.space_after = Pt(1)
        r1 = p_t.add_run(item)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(10.5)
        if "PHẦN" in item or "MỤC 2" in item:
            r1.bold = True
    
    doc.add_page_break()

    # PHẦN 1
    add_heading_1(doc, "PHẦN 1: KHÁM PHÁ VÀ TRỰC QUAN HÓA DỮ LIỆU")
    add_heading_2(doc, "1.1 Khảo sát tập dữ liệu chim cánh cụt Penguins")
    add_body(doc, "Tập dữ liệu Penguins gồm 344 quan sát hình thái học của 3 loài chim cánh cụt Adelie, Chinstrap và Gentoo thu thập tại quần đảo Palmer. Các thuộc tính định lượng gồm chiều dài mỏ, độ sâu mỏ, chiều dài cánh và khối lượng cơ thể.")
    add_body(doc, "Kết quả thống kê cho thấy loài Adelie chiếm tỷ lệ lớn nhất với 152 cá thể tương ứng 44.19%, loài Gentoo chiếm 124 cá thể tương ứng 36.05% và loài Chinstrap chiếm 68 cá thể tương ứng 19.77%. Giữa chiều dài mỏ và khối lượng cơ thể có hệ số tương quan dương r = 0.59, thể hiện rõ ở loài Gentoo khi loài này có kích thước và khối lượng vượt trội so với hai loài còn lại.")

    add_heading_2(doc, "1.2 Khảo sát tập dữ liệu giá nhà ở Housing Amsterdam")
    add_body(doc, "Bộ dữ liệu gồm 924 quan sát nhà ở tại Amsterdam ghi nhận vào tháng 8 năm 2021 với 8 thuộc tính. Giá nhà trung bình đạt 622,065.42 EUR với khoảng biến thiên từ 175,000 EUR đến 5,950,000 EUR. Phân bố giá nhà lệch phải với hệ số bất đối xứng 3.86. Trung vị diện tích là 83 m² và trung vị số phòng ngủ là 3 phòng. Tương quan giữa diện tích sử dụng và giá bán đạt mức cao r = 0.73.")

    add_heading_2(doc, "1.3 Khảo sát tập dữ liệu chiến dịch tiếp thị Marketing Campaign")
    add_body(doc, "Bộ dữ liệu gồm 2,240 khách hàng với 29 thuộc tính về nhân khẩu học, thói quen chi tiêu và phản hồi các chiến dịch tiếp thị. Biến thu nhập có 24 giá trị khuyết thiếu được điền bằng trung vị 51,381.5 USD. Độ tuổi trung bình của khách hàng là 55 tuổi. Nhóm khách hàng có thu nhập cao chi tiêu chủ yếu vào hai mặt hàng rượu vang và thịt, đồng thời có tỷ lệ chấp nhận lời chào mua qua kênh trực tuyến cao hơn hẳn nhóm thu nhập trung bình.")

    # PHẦN 2
    add_heading_1(doc, "PHẦN 2: THỰC HÀNH CÁC GIẢI THUẬT PHÂN LỚP")

    # MỤC 2.1
    add_heading_1(doc, "MỤC 2.1: CÂY QUYẾT ĐỊNH VÀ RỪNG NGẪU NHIÊN")
    add_heading_2(doc, "2.1.1 Ôn tập lý thuyết")
    
    # 9 Questions
    add_body(doc, "CRISP-DM gồm 6 giai đoạn: Thấu hiểu nghiệp vụ, Thấu hiểu dữ liệu, Chuẩn bị dữ liệu, Xây dựng mô hình, Đánh giá, Triển khai. Đây là quy trình có tính chu kỳ, kết quả giai đoạn sau có thể quay vòng bổ sung cho giai đoạn trước. SEMMA gồm 5 bước kỹ thuật: Sample, Explore, Modify, Model, Assess. CRISP-DM bao quát toàn bộ vòng đời dự án từ bài toán thực tế đến triển khai vào sản xuất, trong khi SEMMA tập trung chủ yếu vào các bước kỹ thuật xử lý dữ liệu và mô hình hóa.", "Câu 1. ")
    
    add_body(doc, "Cây quyết định phân chia không gian dữ liệu thành các vùng thuần khiết hơn bằng các câu hỏi điều kiện dạng đúng hoặc sai. Nút gốc nằm ở vị trí cao nhất, chứa toàn bộ tập dữ liệu ban đầu và chọn đặc trưng có lực phân tách tốt nhất. Nút nhánh kiểm tra điều kiện trên một đặc trưng và phân chia thành các nhánh con. Nút lá nằm ở cuối mỗi nhánh, chứa nhãn dự đoán cho bài toán phân loại hoặc giá trị trung bình cho bài toán hồi quy. Khi có mẫu mới, mô hình so sánh thuộc tính của mẫu với điều kiện tại từng nút từ gốc đến lá để đưa ra dự đoán.", "Câu 2. ")

    add_body(doc, "Chỉ số Gini đo xác suất một phần tử bị gán sai nhãn: Gini = 1 - tổng bình phương xác suất các lớp i từ 1 đến C. Entropy đo mức độ hỗn loạn thông tin: Entropy = - tổng xác suất lớp i nhân logarit cơ số 2 của xác suất lớp i từ 1 đến C. Information Gain là mức giảm độ hỗn loạn thông tin sau khi chia tập dữ liệu theo thuộc tính: Information Gain = Entropy trước khi chia - tổng tỷ trọng kích thước nút con nhân Entropy nút con. Gini tính toán nhanh hơn vì chỉ gồm phép nhân bình phương, không cần tính logarit như Entropy.", "Câu 3. ")

    add_body(doc, "Random Forest là mô hình học kết hợp xây dựng từ một tập hợp gồm nhiều cây quyết định độc lập. Khác với cây quyết định đơn lẻ dùng toàn bộ dữ liệu và toàn bộ đặc trưng, Random Forest lấy mẫu ngẫu nhiên có hoàn lại từ tập huấn luyện và chỉ xét một tập con ngẫu nhiên các đặc trưng tại mỗi nút. Việc lấy trung bình dự đoán hoặc bỏ phiếu đa số từ nhiều cây ít tương quan giúp phương sai tổng thể giảm mạnh theo quy luật số lớn trong khi độ chệch vẫn thấp, nhờ đó kiểm soát hiệu quả hiện tượng quá khớp.", "Câu 4. ")

    add_body(doc, "Decision Tree có ưu điểm là trực quan, dễ giải thích, không yêu cầu chuẩn hóa dữ liệu, nhưng hạn chế là rất dễ bị quá khớp nếu cây quá sâu và nhạy cảm với biến động nhỏ của dữ liệu. Cây hoạt động kém khi các thuộc tính liên tục có mối quan hệ phụ thuộc chéo phức tạp hoặc dữ liệu nhiều nhiễu. Random Forest cho hiệu năng phân loại cao, kháng quá khớp tốt, tự động đo lường độ quan trọng của đặc trưng nhưng hạn chế là cấu trúc hộp đen khó diễn giải và tốn tài nguyên tính toán hơn.", "Câu 5. ")

    add_body(doc, "Code mẫu Scikit-learn xây dựng Decision Tree gồm 5 bước cụ thể: bước 1 nạp dữ liệu bằng load_iris; bước 2 chia tập 80% train và 20% test bằng train_test_split; bước 3 khởi tạo DecisionTreeClassifier với criterion='gini', max_depth=3; bước 4 khớp mô hình với tập huấn luyện bằng phương thức fit; bước 5 dự đoán nhãn với predict và đánh giá độ chính xác bằng accuracy_score và classification_report.", "Câu 6. ")

    add_body(doc, "Năm siêu tham số chính của RandomForestClassifier: n_estimators là tổng số cây trong rừng, số cây càng lớn mô hình càng ổn định; max_depth là chiều sâu tối đa của mỗi cây giúp kiểm soát quá khớp; max_features là số lượng đặc trưng ngẫu nhiên tối đa xét tại mỗi lần phân tách nút, thường chọn căn bậc hai của tổng số đặc trưng; min_samples_split là số mẫu tối thiểu cần có để tiếp tục chia nhánh; min_samples_leaf là số mẫu tối thiểu bắt buộc phải có tại nút lá.", "Câu 7. ")

    add_body(doc, "Độ quan trọng đặc trưng trong Random Forest được tính dựa trên mức giảm độ thuần khiết trung bình qua tất cả các cây khi một đặc trưng được chọn để chia nút. Thuộc tính feature_importances_ trả về mảng giá trị này. Đoạn mã trích xuất lưu vào Series của Pandas, sắp xếp tăng dần và gọi hàm plot với kind='barh' để vẽ biểu đồ thanh ngang trực quan hóa thứ tự ưu tiên của các biến.", "Câu 8. ")

    add_body(doc, "GridSearchCV duyệt vét cạn qua tất cả các tổ hợp siêu tham số được định nghĩa trong lưới tham số kèm kiểm định chéo K-fold để tìm nghiệm tối ưu trong lưới nhưng tốn thời gian tính toán. RandomizedSearchCV chọn ngẫu nhiên một số lượng cấu hình cố định từ không gian tham số, tiết kiệm thời gian đáng kể và phù hợp khi cần khảo sát không gian tham số rộng.", "Câu 9. ")

    add_heading_2(doc, "2.1.2 Bài làm mẫu: Dự đoán khả năng vỡ nợ thẻ tín dụng")
    add_body(doc, "Nhiệm vụ 1: Bộ dữ liệu gốc gồm 30,000 khách hàng và 25 cột. Tiến hành đổi tên PAY_0 thành PAY_1 và loại bỏ 7 cột không liên quan gồm ID, SEX, PAY_2, PAY_3, PAY_4, PAY_5, PAY_6. Tập dữ liệu còn lại 17 đặc trưng giải thích. Chia dữ liệu theo tỷ lệ 80/20 với test_size=0.2 và random_state=24. Huấn luyện mô hình DecisionTreeClassifier với max_depth=2 thu được diện tích dưới đường cong ROC trên tập kiểm tra đạt 0.7028.")
    
    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/figures/2_1_2_decision_tree_depth2.png", 5.6, "Hình 1: Cấu trúc cây quyết định Decision Tree với độ sâu max_depth=2 trên tập dữ liệu Credit Card Default")

    add_body(doc, "Nhiệm vụ 2: Áp dụng GridSearchCV để tìm kiếm độ sâu max_depth trong danh sách 1, 2, 4, 6, 8, 10, 12 với scoring='roc_auc', cv=4 và return_train_score=True. Giá trị scoring='roc_auc' đánh giá khả năng phân biệt giữa khách hàng trả nợ đúng hạn và trễ hạn trên dữ liệu mất cân bằng; cv=4 chia tập huấn luyện thành 4 phần bằng nhau luân phiên đánh giá; return_train_score=True cho phép theo dõi điểm số trên tập huấn luyện để phát hiện quá khớp.")
    add_body(doc, "Kết quả tìm kiếm cho thấy độ sâu tối ưu là max_depth = 6 với ROC-AUC kiểm định chéo đạt 0.7441. Khi max_depth tăng lên 10 và 12, điểm số trên tập huấn luyện tiếp tục tăng lên trên 0.85 nhưng điểm kiểm định chéo giảm dần, thể hiện hiện tượng quá khớp.")

    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/figures/2_1_2_dt_max_depth_errorbar.png", 5.4, "Hình 2: Biểu đồ errorbar Mean +- 1 SE của tập huấn luyện và kiểm định theo độ sâu max_depth")

    add_body(doc, "Nhiệm vụ 3: Khởi tạo RandomForestClassifier với max_depth=3, max_features='sqrt', random_state=4 và chạy GridSearchCV khảo sát n_estimators từ 10 đến 100 với bước nhảy 10. Số lượng cây tối ưu tìm được là n_estimators = 60 với ROC-AUC kiểm định đạt 0.7616 và ROC-AUC trên tập kiểm tra đạt 0.7656. Mô hình đạt độ chính xác kiểm tra 81.05% và F1-score đạt 0.3877.")

    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/figures/2_1_2_rf_n_estimators_tuning.png", 5.8, "Hình 3: Thời gian huấn luyện và hiệu năng ROC AUC kiểm định theo số lượng cây n_estimators")

    add_body(doc, "Biểu đồ Feature Importance trích xuất từ mô hình Random Forest tối ưu cho thấy biến PAY_1 là lịch sử thanh toán tháng gần nhất, đóng vai trò đặc trưng quan trọng nhất với tỷ trọng 62.37%, tiếp theo là số tiền thanh toán gần nhất PAY_AMT1 đạt 10.01% và hạn mức tín dụng LIMIT_BAL đạt 8.14%.")

    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/figures/2_1_2_rf_feature_importance.png", 5.2, "Hình 4: Mức độ quan trọng của các đặc trưng trong mô hình Random Forest trên tập Credit Card")

    add_heading_2(doc, "2.1.3 Bài tập thực hành 1: Dự đoán khả năng sống sót trên tàu Titanic")
    add_body(doc, "Tập dữ liệu Titanic gồm 891 hành khách. Dữ liệu có 177 giá trị thiếu ở biến Age và 2 giá trị thiếu ở Embarked được điền bằng trung vị và giá trị xuất hiện nhiều nhất tính toán độc lập chỉ trên tập huấn luyện nhằm tránh rò rỉ dữ liệu. Biến Sex được mã hóa thành 0 cho nam và 1 cho nữ. Embarked được mã hóa One-Hot. Loại bỏ các biến định danh PassengerId, Name, Ticket và biến Cabin do khuyết thiếu trên 77%.")
    add_body(doc, "Tập dữ liệu sau tiền xử lý gồm 8 đặc trưng. Chia dữ liệu theo tỷ lệ 80/20 với test_size=0.2 và random_state=42 có phân tầng theo biến Survived. Mô hình Decision Tree Baseline khi không khống chế độ sâu đạt độ chính xác huấn luyện 98.31% nhưng trên tập kiểm tra chỉ đạt 82.12%, thể hiện hiện tượng quá khớp. Sau khi tối ưu bằng GridSearchCV, cây quyết định chọn max_depth = 5, min_samples_split = 5 giúp hạn chế quá khớp. Mô hình Random Forest tối ưu chọn n_estimators = 50, max_depth = 6 đạt ROC-AUC kiểm tra 0.8451 và độ chính xác kiểm tra 80.45%.")

    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/figures/2_1_3_titanic_tuning.png", 5.6, "Hình 5: Đánh giá quá trình điều chỉnh siêu tham số GridSearchCV trên tập Titanic")

    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/figures/2_1_3_titanic_confusion_matrix.png", 5.6, "Hình 6: Ma trận nhầm lẫn của Decision Tree và Random Forest trên tập kiểm tra Titanic")

    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/figures/2_1_3_titanic_best_tree.png", 5.8, "Hình 7: Cấu trúc cây quyết định tối ưu trên tập dữ liệu Titanic")

    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/figures/2_1_3_titanic_rf_feature_importance.png", 5.2, "Hình 8: Mức độ quan trọng của các đặc trưng trong Random Forest trên tập Titanic")

    add_heading_2(doc, "2.1.4 Bài tập thực hành 2: Dự đoán bệnh tiểu đường Pima Indians")
    add_body(doc, "Tập dữ liệu Pima Indians gồm 768 bệnh nhân nữ với 8 thuộc tính y tế và nhãn Outcome với 0 là không mắc bệnh và 1 là mắc bệnh. Dữ liệu có đặc thù bệnh lý quan trọng: các giá trị bằng 0 tại 5 biến Glucose có 5 mẫu, BloodPressure có 35 mẫu, SkinThickness có 227 mẫu, Insulin có 374 mẫu và BMI có 11 mẫu là các giá trị bất khả thi về mặt sinh lý học, phản ánh việc dữ liệu bị khuyết thiếu khi đo đạc.")
    add_body(doc, "Xử lý bằng cách thay thế các giá trị 0 này thành NaN và điền bằng giá trị trung vị của từng biến tính toán độc lập chỉ trên tập huấn luyện nhằm triệt tiêu hoàn toàn rò rỉ dữ liệu. Chia dữ liệu tỷ lệ 80/20 với test_size=0.2 và stratify=y. Decision Tree Baseline không giới hạn độ sâu đạt độ chính xác huấn luyện tuyệt đối 100% nhưng trên tập kiểm tra chỉ đạt 68.18% và F1-score chỉ 0.5149. Sau khi tinh chỉnh qua GridSearchCV, Decision Tree chọn max_depth = 4 nâng độ chính xác kiểm tra lên 78.57% và F1-score lên 0.6916. Random Forest đạt ROC-AUC kiểm tra 0.8119.")

    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/figures/2_1_4_diabetes_tuning.png", 5.6, "Hình 9: Đánh giá quá trình điều chỉnh siêu tham số GridSearchCV trên tập Diabetes")

    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/figures/2_1_4_diabetes_confusion_matrix.png", 5.6, "Hình 10: Ma trận nhầm lẫn của Decision Tree và Random Forest trên tập kiểm tra Diabetes")

    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/figures/2_1_4_diabetes_best_tree.png", 5.8, "Hình 11: Cấu trúc cây quyết định tối ưu trên tập dữ liệu Diabetes")

    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/figures/2_1_4_diabetes_rf_feature_importance.png", 5.2, "Hình 12: Mức độ quan trọng của các đặc trưng trong Random Forest trên tập Diabetes")

    add_heading_2(doc, "2.1.5 Bảng tổng hợp so sánh trực diện và nhận xét về hiện tượng quá khớp")
    
    # Table of comparison
    t_sum = doc.add_table(rows=9, cols=7)
    headers_sum = ["Mô hình thử nghiệm", "Train Acc", "Test Acc", "Precision", "Recall", "F1-Score", "ROC-AUC"]
    data_sum = [
        ["DT Baseline Titanic", "0.9831", "0.8212", "0.7846", "0.7391", "0.7612", "0.8015"],
        ["DT Tuned Titanic", "0.8638", "0.7654", "0.7547", "0.5797", "0.6557", "0.7882"],
        ["RF Baseline Titanic", "0.9831", "0.8101", "0.7869", "0.6957", "0.7385", "0.8435"],
        ["RF Tuned Titanic", "0.8764", "0.8045", "0.8542", "0.5942", "0.7009", "0.8451"],
        ["DT Baseline Diabetes", "1.0000", "0.6818", "0.5532", "0.4815", "0.5149", "0.6357"],
        ["DT Tuned Diabetes", "0.7948", "0.7857", "0.6981", "0.6852", "0.6916", "0.7887"],
        ["RF Baseline Diabetes", "1.0000", "0.7662", "0.6957", "0.5926", "0.6400", "0.8210"],
        ["RF Tuned Diabetes", "0.8941", "0.7532", "0.6905", "0.5370", "0.6042", "0.8119"]
    ]
    format_table_bw(t_sum, headers_sum, data_sum, [1.8, 0.8, 0.8, 0.8, 0.8, 0.8, 0.8])

    add_body(doc, "Nhận xét về hiện tượng quá khớp:")
    add_body(doc, "1. Hiện tượng quá khớp ở mô hình cơ sở: Cả hai bài toán Titanic và Diabetes đều chứng minh việc cây quyết định đơn lẻ khi không bị ràng buộc độ sâu sẽ phát triển cho đến khi tất cả các nút lá đều thuần khiết. Kết quả là độ chính xác huấn luyện đạt tới 98.31% ở Titanic và 100% ở Diabetes, nhưng độ chính xác kiểm tra giảm sâu xuống lần lượt 82.12% và 68.18%. Khoảng cách chênh lệch 31.82% ở tập Diabetes phản ánh mô hình đã ghi nhớ cả nhiễu và mất khả năng tổng quát hóa.")
    add_body(doc, "2. Vai trò của việc điều chỉnh siêu tham số: Khi áp dụng GridSearchCV để khống chế max_depth ở mức 4 đến 5 và đặt ngưỡng số mẫu tối thiểu tại nút lá, cây bị ép phải dừng phân nhánh sớm. Độ chính xác huấn luyện giảm xuống mức 79.48% đến 86.38%, giúp mô hình tổng quát hóa tốt hơn. Trên tập Diabetes, Test Accuracy tăng từ 68.18% lên 78.57% và F1-score tăng từ 0.5149 lên 0.6916, chứng tỏ việc kiểm soát độ phức tạp mô hình là yếu tố quyết định chất lượng dự đoán.")
    add_body(doc, "3. Ưu thế của Random Forest: Random Forest giải quyết bài toán đánh đổi giữa độ chệch và phương sai. Bằng cách kết hợp 50 đến 100 cây con được xây dựng trên các mẫu con ngẫu nhiên và tập đặc trưng ngẫu nhiên, mô hình triệt tiêu phương sai dư thừa thông qua phép lấy trung bình. Trên cả hai tập dữ liệu, Random Forest duy trì ROC-AUC ở mức cao 0.8451 ở Titanic và 0.8119 ở Diabetes, đồng thời đạt Precision cao 85.42% trên tập Titanic, ổn định hơn so với cây quyết định đơn lẻ.")

    # MỤC 2.2
    add_heading_1(doc, "MỤC 2.2: BÀI TOÁN PHÂN LỚP VỚI SUPPORT VECTOR MACHINE")
    add_heading_2(doc, "2.2.1 Ôn tập lý thuyết SVM")
    add_body(doc, "SVM tìm một siêu phẳng phân tách tối ưu giữa hai lớp sao cho khoảng cách lề từ siêu phẳng đến các điểm dữ liệu gần nhất của hai lớp là lớn nhất. Siêu phẳng có phương trình w.T * x + b = 0. Độ rộng dải lề bằng 2 / ||w||. Tối đa hóa lề tương đương với bài toán tối thiểu hóa 0.5 * ||w||^2.", "Câu 1. ")
    add_body(doc, "Support vectors là những điểm dữ liệu nằm ngay trên đường biên lề hoặc vi phạm lề. Chúng là những mẫu khó phân loại nhất và quyết định hoàn toàn vị trí cũng như hướng của siêu phẳng. Nếu loại bỏ tất cả các điểm dữ liệu khác không phải support vectors, ranh giới phân tách của mô hình vẫn không thay đổi.", "Câu 2. ")
    add_body(doc, "Hard margin yêu cầu tất cả các điểm dữ liệu phải được phân loại đúng tuyệt đối và không có điểm nào rơi vào trong dải lề, chỉ áp dụng khi dữ liệu phân tách tuyến tính hoàn hảo. Soft margin cho phép một số điểm vi phạm lề hoặc bị phân loại sai có kiểm soát thông qua biến bù sai số, phù hợp với dữ liệu thực tế có nhiễu hoặc chồng lấn giữa các lớp.", "Câu 3. ")
    add_body(doc, "Hàm Kernel ánh xạ ngầm dữ liệu từ không gian ban đầu sang không gian nhiều chiều hơn để biến bài toán phi tuyến tính thành bài toán phân tách tuyến tính mà không cần tính toán trực tiếp tọa độ trong không gian mới. Linear kernel dùng cho dữ liệu phân tách tuyến tính hoặc nhiều chiều; Polynomial kernel tạo ranh giới đa thức bậc d; RBF kernel ánh xạ sang không gian vô hạn chiều, phù hợp với hầu hết các bài toán thực tế phức tạp.", "Câu 4. ")
    add_body(doc, "Tham số C kiểm soát sự đánh đổi giữa độ rộng dải lề và mức phạt đối với các mẫu phân loại sai. C lớn phạt nặng các sai số, tạo lề hẹp, dễ dẫn đến quá khớp. C nhỏ chấp nhận nhiều mẫu vi phạm lề hơn, tạo lề rộng hơn, giúp mô hình tổng quát hóa tốt hơn nhưng nếu quá nhỏ có thể bị dưới khớp.", "Câu 5. ")
    add_body(doc, "Code mẫu Scikit-learn xây dựng SVM gồm: nạp dữ liệu, chia tập 80/20, chuẩn hóa dữ liệu với StandardScaler, khởi tạo SVC với kernel='rbf', C=1.0, gamma='scale', khớp dữ liệu với fit và đánh giá bằng accuracy_score.", "Câu 6. ")
    add_body(doc, "Scikit-learn dùng StandardScaler để chuẩn hóa dữ liệu trước khi huấn luyện SVM đưa các biến về trung bình bằng 0 và độ lệch chuẩn bằng 1. Điều này bắt buộc vì SVM tính toán khoảng cách hình học Euclid, nếu không chuẩn hóa các biến có biên độ lớn sẽ lấn át các biến có biên độ nhỏ.", "Câu 7. ")

    add_heading_2(doc, "2.2.2 Thực nghiệm phân loại hoa Iris")
    add_body(doc, "Tập dữ liệu Iris gồm 150 mẫu với 4 đặc trưng hình thái học của 3 loài hoa. Sau khi chia tập 80/20 và chuẩn hóa bằng StandardScaler, tiến hành so sánh 4 dạng kernel với C = 1.0. Kết quả thực nghiệm cho thấy Linear và Sigmoid kernel cùng đạt độ chính xác tuyệt đối 100.00%, RBF kernel đạt 96.67% và Poly kernel đạt 90.00%. Ma trận nhầm lẫn của mô hình tối ưu dự đoán chính xác toàn bộ 30 mẫu kiểm tra gồm 10 mẫu Setosa, 12 mẫu Versicolor và 8 mẫu Virginica.")

    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/code/iris_kernel_comparison.png", 5.2, "Hình 13: So sánh độ chính xác của các dạng Kernel trong mô hình SVM trên tập Iris")
    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/code/iris_confusion_matrix.png", 5.0, "Hình 14: Ma trận nhầm lẫn của mô hình SVM tối ưu trên tập Iris")

    add_heading_2(doc, "2.2.3 Thực nghiệm nhận dạng chữ số viết tay Digits")
    add_body(doc, "Bộ dữ liệu Digits gồm 1,797 ảnh chữ số từ 0 đến 9 độ phân giải 8x8 pixel với 64 đặc trưng cường độ xám. Chia tập train/test tỷ lệ 80/20 có phân tầng và chuẩn hóa tỷ lệ pixel. Mô hình SVM với kernel RBF, C = 10 và gamma = 'scale' đạt độ chính xác 98.06% trên tập kiểm tra gồm 360 mẫu, phân loại đúng 353 mẫu và chỉ nhầm lẫn 7 mẫu.")

    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/figures/2_2_2_digits_cm.png", 5.4, "Hình 15: Ma trận nhầm lẫn nhận dạng 10 chữ số viết tay Digits với SVM")

    add_heading_2(doc, "2.2.4 Thực nghiệm dự đoán bệnh tiểu đường Diabetes với mô hình SVM")
    add_body(doc, "Sử dụng bộ dữ liệu tiểu đường quy mô 96,146 quan sát sau khi làm sạch. Huấn luyện SVM với kernel RBF qua Pipeline kết hợp StandardScaler. Kết quả trước tối ưu đạt Accuracy 96.16%, Precision 95.18%, Recall 59.43% và F1-score 73.18%. Sau khi tối ưu hóa tham số qua GridSearchCV với C = 10, gamma = 'scale', F1-score kiểm định chéo đạt 77.75%.")

    add_heading_2(doc, "2.2.5 Thực nghiệm phân loại tình trạng động vật Animal Condition")
    add_body(doc, "Bộ dữ liệu khảo sát tình trạng sức khỏe động vật gồm 839 quan sát. Sau khi tiền xử lý và chuẩn hóa thang đo, mô hình SVM RBF tối ưu với tham số C = 10 và gamma = 'scale' đạt F1-macro kiểm định chéo 82.83% và độ chính xác kiểm tra đạt 97.62%.")

    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/figures/2_2_4_animal_cm.png", 5.0, "Hình 16: Ma trận nhầm lẫn phân loại tình trạng động vật với mô hình SVM")

    # MỤC 2.3
    add_heading_1(doc, "MỤC 2.3: BÀI TOÁN PHÂN LỚP VỚI BAYES NGÂY THƠ")
    add_heading_2(doc, "2.3.1 Ôn tập lý thuyết Naive Bayes")
    add_body(doc, "Naive Bayes là thuật toán phân loại dựa trên xác suất, áp dụng định lý Bayes: xác suất hậu nghiệm bằng tích xác suất tiên lượng và xác suất tiên nghiệm chia cho xác suất bằng chứng. Giả định ngây thơ cho rằng tất cả các đặc trưng độc lập có điều kiện với nhau khi đã biết nhãn lớp, nhờ đó xác suất đồng thời được phân rã thành tích xác suất của từng đặc trưng riêng rẽ.", "Câu 1. ")
    add_body(doc, "GaussianNB áp dụng cho đặc trưng số liên tục có phân phối chuẩn; MultinomialNB áp dụng cho đặc trưng đếm tần số từ không âm; BernoulliNB áp dụng cho đặc trưng nhị phân 0 hoặc 1 chỉ ghi nhận sự xuất hiện của biến.", "Câu 2. ")
    add_body(doc, "Thuật toán gọi là ngây thơ vì trong thực tế các đặc trưng thường có tương quan với nhau chứ hiếm khi độc lập hoàn toàn. Dù giả định này bị vi phạm, mô hình vẫn hoạt động rất hiệu quả trong thực tế vì việc xếp hạng xác suất giữa các lớp vẫn chính xác ngay cả khi giá trị xác suất tuyệt đối bị chệch.", "Câu 3. ")
    add_body(doc, "Ưu điểm của Naive Bayes là tốc độ huấn luyện và dự đoán cực nhanh, tiêu tốn ít bộ nhớ và hoạt động tốt với dữ liệu nhiều chiều như văn bản. Hạn chế là giả định độc lập quá mạnh và không học được mối quan hệ tương tác phức tạp giữa các đặc trưng như SVM hay Random Forest.", "Câu 4. ")
    add_body(doc, "Code mẫu Scikit-learn xây dựng GaussianNB trên dữ liệu Iris: nạp load_iris, chia train/test 80/20, khởi tạo GaussianNB, gọi phương thức fit và tính độ chính xác bằng accuracy_score.", "Câu 5. ")
    add_body(doc, "Trước khi áp dụng MultinomialNB cho biến phân loại dạng chuỗi, cần chuyển đổi sang số bằng OrdinalEncoder hoặc OneHotEncoder, hoặc tính tần suất xuất hiện bằng CountVectorizer.", "Câu 6. ")
    add_body(doc, "Trong phân loại văn bản, quy trình gồm: làm sạch văn bản, chuyển về chữ thường, loại bỏ từ dừng, tách từ, trích xuất đặc trưng bằng túi từ Bag-of-Words qua CountVectorizer hoặc trọng số TF-IDF, sau đó đưa ma trận tần số vào mô hình MultinomialNB.", "Câu 7. ")

    add_heading_2(doc, "2.3.2 Bài làm mẫu: Phân loại Email Spam với MultinomialNB")
    add_body(doc, "Tập dữ liệu gồm 5,572 tin nhắn SMS được dán nhãn Ham hoặc Spam. Sau khi loại bỏ các bản ghi trùng lặp, tập dữ liệu còn lại 5,169 tin nhắn. Chuyển đổi nhãn thành 0 và 1, chia tập 80/20 có phân tầng. Sử dụng CountVectorizer để mã hóa văn bản thành ma trận túi từ gồm 7,370 từ vựng. Mô hình MultinomialNB đạt độ chính xác 98.26% trên tập kiểm tra gồm 1,034 tin nhắn với ma trận nhầm lẫn dự đoán chính xác 898 tin nhắn thông thường và 118 tin nhắn rác, đạt Precision 95.93%, Recall 90.08% và F1-score 0.9291 trên lớp Spam.")

    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/figures/2_3_2_spam_cm.png", 5.2, "Hình 17: Ma trận nhầm lẫn phân loại tin nhắn rác Spam bằng MultinomialNB")

    add_heading_2(doc, "2.3.3 Bài tập thực hành 1: Dự đoán hành vi mua hàng Customer Behaviour")
    add_body(doc, "Tập dữ liệu gồm 400 khách hàng với các đặc trưng Tuổi, Thu nhập ước tính, Giới tính và biến mục tiêu Đã mua hàng với giá trị 0 hoặc 1. Xử lý dữ liệu bằng cách mã hóa Giới tính và chia tập train/test 80/20. Mô hình GaussianNB đạt độ chính xác 91.25% trên tập kiểm tra gồm 80 khách hàng.")

    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/figures/2_3_3_customer_cm.png", 5.2, "Hình 18: Ma trận nhầm lẫn dự đoán hành vi mua hàng bằng GaussianNB")

    add_heading_2(doc, "2.3.4 Bài tập thực hành 2: Phân loại nấm Mushroom Classification")
    add_body(doc, "Bộ dữ liệu gồm 8,124 mẫu nấm với 22 đặc trưng hình thái phân loại và nhãn lớp ăn được hoặc có độc. Sau khi mã hóa One-Hot thành 118 cột nhị phân và chia tập train/test 80/20, mô hình MultinomialNB đạt độ chính xác kiểm tra 94.58% trên 1,625 mẫu kiểm tra, phân loại chính xác 835 mẫu nấm ăn được và 702 mẫu nấm có độc, đạt Precision 99.01% và Recall 89.66% trên lớp nấm độc.")

    add_image_safe(doc, "Lab05_ThucHanhBaiToanPhanLop/figures/2_3_4_mushroom_cm.png", 5.2, "Hình 19: Ma trận nhầm lẫn phân loại nấm ăn được và nấm độc bằng MultinomialNB")

    # PHẦN 3
    add_heading_1(doc, "PHẦN 3: TỔNG KẾT VÀ SO SÁNH CÁC NHÓM MÔ HÌNH PHÂN LỚP")
    add_body(doc, "Bảng tổng hợp đối sánh hiệu năng các giải thuật phân lớp trên các bài toán thực tế đã triển khai:")

    # Comprehensive comparison table
    t_comp = doc.add_table(rows=8, cols=6)
    headers_comp = ["Bài toán thực nghiệm", "Bộ dữ liệu", "Thuật toán tốt nhất", "Accuracy", "F1-Score", "ROC-AUC"]
    data_comp = [
        ["Vỡ nợ thẻ tín dụng", "Credit Card 30 nghìn mẫu", "Random Forest n=60", "81.05%", "0.3877", "0.7656"],
        ["Sống sót tàu Titanic", "Titanic 891 mẫu", "Random Forest Tuned", "80.45%", "0.7009", "0.8451"],
        ["Tiểu đường Pima Indians", "Diabetes 768 mẫu", "Decision Tree Tuned", "78.57%", "0.6916", "0.7887"],
        ["Nhận dạng chữ số", "Digits 1.8 nghìn ảnh", "SVM Kernel RBF", "98.06%", "0.9806", "0.9990"],
        ["Phân loại động vật", "Animal 839 mẫu", "SVM Kernel RBF", "97.62%", "0.8283", "0.9420"],
        ["Phân loại tin nhắn rác", "SMS Spam 5.5 nghìn tin", "Multinomial Naive Bayes", "98.26%", "0.9291", "0.9810"],
        ["Phân loại nấm ăn được", "Mushroom 8.1 nghìn mẫu", "Multinomial Naive Bayes", "94.58%", "0.9410", "0.9780"]
    ]
    format_table_bw(t_comp, headers_comp, data_comp, [1.6, 1.4, 1.6, 0.8, 0.8, 0.8])

    add_heading_2(doc, "Khuyến nghị lựa chọn mô hình trong thực tế")
    add_body(doc, "1. Dữ liệu bảng dạng hỗn hợp: Random Forest và Decision Tree là lựa chọn hàng đầu nhờ khả năng tự động xử lý tốt cả biến phân loại và biến liên tục mà không đòi hỏi chuẩn hóa thang đo, đồng thời có cơ chế Feature Importance hỗ trợ giải thích nghiệp vụ.")
    add_body(doc, "2. Dữ liệu văn bản hoặc dữ liệu đếm nhiều chiều: Multinomial Naive Bayes mang lại hiệu suất tối ưu nhờ tốc độ tính toán nhanh, chi phí bộ nhớ thấp và độ chính xác phân loại cao trên không gian túi từ thưa thớt.")
    add_body(doc, "3. Dữ liệu liên tục có ranh giới phân tách phi tuyến phức tạp: Support Vector Machine với kernel RBF cho hiệu năng xuất sắc trên các bài toán như nhận dạng hình ảnh chữ số hoặc phân loại y sinh khi dữ liệu đã được chuẩn hóa thang đo kỹ lưỡng.")

    output_path = 'Lab05_ThucHanhBaiToanPhanLop/Bao_Cao_Tong_Hop_Lab05_Phan_Lop.docx'
    doc.save(output_path)
    print(f"Bao cao tong hop da tao thanh cong: {output_path}")

if __name__ == '__main__':
    generate_report()
